import io
import json
import os
import sys
import tempfile
import unittest

from docx import Document
from docx.shared import Cm, Pt
from pptx import Presentation


PROJECT_FOLDER = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BACKEND_FOLDER = os.path.join(PROJECT_FOLDER, "backend")
sys.path.insert(0, BACKEND_FOLDER)

os.environ["OFORMITEL_DATA_FOLDER"] = tempfile.mkdtemp(prefix="oformitel-tests-")

from app import OUTPUTS_FOLDER, PROFILES_FOLDER, app  # noqa: E402
from engine import analyze_layout, fix_toc  # noqa: E402
from pptx_engine import analyze_presentation  # noqa: E402


def make_docx():
    file_data = io.BytesIO()
    document = Document()
    document.add_paragraph("Проверочный документ")
    document.add_paragraph("Текст для проверки обработки файла.")
    document.save(file_data)
    file_data.seek(0)
    return file_data


def make_badly_formatted_docx():
    file_data = io.BytesIO()
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1)
    section.bottom_margin = Cm(1)
    section.left_margin = Cm(1)
    section.right_margin = Cm(1)
    paragraph = document.add_paragraph("Документ для проверки пользовательского профиля.")
    run = paragraph.runs[0]
    run.font.name = "Arial"
    run.font.size = Pt(10)
    paragraph.paragraph_format.line_spacing = 1
    paragraph.paragraph_format.first_line_indent = Cm(0)
    document.save(file_data)
    file_data.seek(0)
    return file_data


def make_pptx():
    file_data = io.BytesIO()
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = "Проверочный слайд."
    presentation.save(file_data)
    file_data.seek(0)
    return file_data


def make_pptx_with_traps(path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = "Заголовок с точкой."
    slide.placeholders[1].text = "а" * 401
    presentation.save(path)


def make_toc_docx():
    document = Document()
    document.add_paragraph("Оглавление")
    entry = document.add_paragraph("Введение........................................3")
    entry.paragraph_format.left_indent = Cm(0.8)
    entry = document.add_paragraph("Заключение.....................................10")
    entry.paragraph_format.left_indent = Cm(0.8)
    document.add_paragraph()
    document.add_paragraph("Введение")
    document.add_paragraph("Текст введения длиной больше ста символов нужен только для того, чтобы анализатор распознал начало основного текста документа.")
    return document


class SmokeTests(unittest.TestCase):
    def setUp(self):
        self.client = app.test_client()

    def test_main_page_opens(self):
        response = self.client.get("/")

        self.assertEqual(response.status_code, 200)

    def test_title_page_is_created(self):
        response = self.client.post(
            "/create_title_page",
            data={
                "organization": "Тестовая организация",
                "work_title": "Проверочная работа",
                "participants": "Иванов И. И.",
                "supervisor": "Петров П. П.",
                "city": "Москва",
                "year": "2026",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn("file", response.get_json())

    def test_docx_is_processed(self):
        response = self.client.post(
            "/apply",
            data={
                "document": (make_docx(), "test.docx"),
                "profile_file": "01_inzhenery_budushchego.json",
                "skip_title": "1",
            },
            content_type="multipart/form-data",
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn("gost", response.get_json())

    def test_custom_profile_is_saved_and_applied(self):
        profile_values = {
            "name": "Проверочный профиль",
            "font_name": "Arial",
            "font_size_pt": "11",
            "line_spacing": "2",
            "first_line_indent_cm": "0.5",
            "margin_top_cm": "1.1",
            "margin_bottom_cm": "1.2",
            "margin_left_cm": "1.3",
            "margin_right_cm": "1.4",
            "document_type": "research",
        }
        response = self.client.post("/new_profile", data=profile_values)
        self.assertEqual(response.status_code, 302)

        created_profiles = []
        for file_name in os.listdir(PROFILES_FOLDER):
            if not file_name.startswith("profile_") or not file_name.endswith(".json"):
                continue
            path = os.path.join(PROFILES_FOLDER, file_name)
            with open(path, "r", encoding="utf-8") as profile_file:
                profile = json.load(profile_file)
            if profile.get("name") == profile_values["name"]:
                created_profiles.append(file_name)
        self.assertEqual(len(created_profiles), 1)
        with open(os.path.join(PROFILES_FOLDER, created_profiles[0]), "r", encoding="utf-8") as profile_file:
            saved_profile = json.load(profile_file)
        self.assertEqual(saved_profile["document_type"], "research")
        self.assertTrue(saved_profile["structure"]["toc"])

        response = self.client.post(
            "/apply",
            data={
                "document": (make_badly_formatted_docx(), "custom-profile.docx"),
                "profile_file": created_profiles[0],
                "skip_title": "1",
            },
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 200)
        result = response.get_json()
        self.assertIn("url", result)
        self.assertTrue(any("обязательный титульный лист" in warning for warning in result["warnings"]))
        self.assertTrue(any("обязательное оглавление" in warning for warning in result["warnings"]))

        output_name = os.path.basename(result["url"])
        output = Document(os.path.join(OUTPUTS_FOLDER, output_name))
        section = output.sections[0]
        self.assertAlmostEqual(section.top_margin.cm, 1.1, places=1)
        self.assertAlmostEqual(section.bottom_margin.cm, 1.2, places=1)
        self.assertAlmostEqual(section.left_margin.cm, 1.3, places=1)
        self.assertAlmostEqual(section.right_margin.cm, 1.4, places=1)
        paragraph = output.paragraphs[0]
        run = paragraph.runs[0]
        self.assertEqual(run.font.name, "Arial")
        self.assertEqual(run.font.size.pt, 11)
        self.assertAlmostEqual(paragraph.paragraph_format.line_spacing, 2)
        self.assertAlmostEqual(paragraph.paragraph_format.first_line_indent.cm, 0.5, places=1)

    def test_invalid_custom_profile_values_are_rejected(self):
        response = self.client.post(
            "/new_profile",
            data={
                "name": "Невалидный профиль",
                "font_name": "Arial",
                "font_size_pt": "-10",
                "line_spacing": "0",
                "first_line_indent_cm": "-1",
                "margin_top_cm": "-1",
                "margin_bottom_cm": "-1",
                "margin_left_cm": "-1",
                "margin_right_cm": "-1",
            },
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn("Размер шрифта", response.get_data(as_text=True))

    def test_pptx_is_analyzed(self):
        response = self.client.post(
            "/analyze_pptx",
            data={"presentation": (make_pptx(), "test.pptx")},
            content_type="multipart/form-data",
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.get_json()["issues"]), 1)

    def test_empty_title_page_is_rejected(self):
        response = self.client.post("/create_title_page", data={})

        self.assertEqual(response.status_code, 400)

    def test_stylesheet_is_available_to_packaged_app(self):
        response = self.client.get("/style.css")

        self.assertEqual(response.status_code, 200)
        self.assertIn(b"--blue", response.data)

    def test_toc_trap_keeps_page_number_on_same_line(self):
        document = make_toc_docx()
        fix_toc(document)

        layout = analyze_layout(document)
        entries = [entry for entry in layout["toc"] if entry.text.strip()]
        self.assertEqual([entry.text for entry in entries], ["Введение\t3", "Заключение\t10"])
        self.assertTrue(all(entry.paragraph_format.left_indent.cm == 0 for entry in entries))

    def test_presentation_traps_are_reported(self):
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, "traps.pptx")
            make_pptx_with_traps(path)
            issues = analyze_presentation(path)

        issue_types = [issue["issue_type"] for issue in issues]
        self.assertIn("Точка в конце заголовка", issue_types)
        self.assertIn("Слишком длинный текст на слайде", issue_types)


if __name__ == "__main__":
    unittest.main()
