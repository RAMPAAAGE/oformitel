from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.shared import Cm, Pt

FONT_NAME = "Times New Roman"


def add_styled_paragraph(document, text, alignment, bold, font_size_pt, make_uppercase=False):
    paragraph = document.add_paragraph()
    paragraph.alignment = alignment
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    if make_uppercase:
        text = text.upper()

    run = paragraph.add_run(text)
    run.font.name = FONT_NAME
    run.font.size = Pt(font_size_pt)
    run.font.bold = bold

    return paragraph


def add_empty_paragraphs(document, how_many):
    for _ in range(how_many):
        document.add_paragraph()


def check_required_fields(fields):
    required_field_names = [
        "organization",
        "work_title",
        "participants",
        "supervisor",
        "city",
        "year",
    ]

    for field_name in required_field_names:
        if not fields.get(field_name, "").strip():
            raise ValueError("Заполните все поля титульного листа.")


def create_title_page_document(fields, output_path):
    check_required_fields(fields)

    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

    add_styled_paragraph(
        document,
        fields["organization"],
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        bold=True,
        font_size_pt=14,
    )

    add_empty_paragraphs(document, 6)

    add_styled_paragraph(
        document,
        fields["work_title"],
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        bold=True,
        font_size_pt=16,
        make_uppercase=True,
    )

    add_empty_paragraphs(document, 6)

    add_styled_paragraph(
        document,
        fields["participants"],
        alignment=WD_ALIGN_PARAGRAPH.RIGHT,
        bold=False,
        font_size_pt=14,
    )
    add_styled_paragraph(
        document,
        fields["supervisor"],
        alignment=WD_ALIGN_PARAGRAPH.RIGHT,
        bold=False,
        font_size_pt=14,
    )

    add_empty_paragraphs(document, 4)

    city_and_year = fields["city"] + ", " + fields["year"]
    add_styled_paragraph(
        document,
        city_and_year,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        bold=True,
        font_size_pt=14,
    )

    document.save(output_path)
    return output_path
