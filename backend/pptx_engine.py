import os

from pptx import Presentation

# Если текста на слайде больше, чем это число символов, считаем,
# что текста слишком много и его стоит сократить.
MAX_CHARACTERS_PER_SLIDE = 400


def validate_pptx_file(file_path):
    if not file_path or not os.path.isfile(file_path):
        raise ValueError("Файл презентации не найден.")

    if not file_path.lower().endswith(".pptx"):
        raise ValueError("Можно выбрать только файл с расширением .pptx.")

    return True


def get_slide_title(slide):
    # Не у каждого слайда есть заголовок (например, у слайда с одной картинкой).
    if slide.shapes.title is None:
        return ""

    if slide.shapes.title.text is None:
        return ""

    return slide.shapes.title.text


def count_characters_on_slide(slide):
    total_characters = 0

    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue

        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                total_characters += len(run.text)

    return total_characters


def check_title_ends_with_dot(slide_number, title, issues_list):
    title = title.strip()

    if title.endswith("."):
        issues_list.append(
            {
                "slide_number": slide_number,
                "issue_type": "Точка в конце заголовка",
                "details": "Заголовок «" + title + "» заканчивается точкой. По правилам оформления точка в конце заголовка не ставится.",
            }
        )


def check_too_much_text(slide_number, character_count, issues_list):
    if character_count > MAX_CHARACTERS_PER_SLIDE:
        issues_list.append(
            {
                "slide_number": slide_number,
                "issue_type": "Слишком длинный текст на слайде",
                "details": "На слайде " + str(character_count) + " символов текста. Рекомендуется не больше " + str(MAX_CHARACTERS_PER_SLIDE) + ".",
            }
        )


def analyze_presentation(file_path):
    validate_pptx_file(file_path)
    presentation = Presentation(file_path)

    issues_list = []
    slide_number = 0

    for slide in presentation.slides:
        slide_number += 1

        title = get_slide_title(slide)
        check_title_ends_with_dot(slide_number, title, issues_list)

        character_count = count_characters_on_slide(slide)
        check_too_much_text(slide_number, character_count, issues_list)

    return issues_list
