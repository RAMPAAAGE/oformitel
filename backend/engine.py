import copy
import json
import os

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from punctuation import BIBLIO, apply_edits, find_edits, get_nodes, heading_edit, make_changes, spell_words

MAX_DOCUMENT_SIZE_MB = 5

CRITERIA_NAMES = {
    1: "поля страницы",
    2: "шрифт",
    3: "размер шрифта",
    4: "межстрочный интервал",
    5: "отступ первой строки",
    6: "нумерация страниц",
    7: "точки в оглавлении",
    8: "нумерация картинок",
    9: "тема с новой страницы",
    10: "пунктуация",
}

KNOWN_TITLES = [
    "оглавление", "содержание", "введение", "заключение",
    "список литературы", "список использованных источников", "литература",
]
TOC_NAMES = ["оглавление", "содержание"]
TITLE_LIMIT = 40  # титульный лист ищем только в начале
SPACING_ROLES = ["body", "heading", "list", "center"]

DOCUMENT_TYPES = {
    "research": {
        "label": "Исследовательская работа",
        "title_page": True,
        "toc": True,
        "sections": ["введение", "заключение", "список литературы"],
    },
    "referat": {
        "label": "Реферат",
        "title_page": True,
        "toc": True,
        "sections": ["введение", "заключение", "список литературы"],
    },
    "practical": {
        "label": "Практическая работа",
        "title_page": True,
        "toc": False,
        "sections": ["цель", "ход работы", "результат", "заключение"],
    },
    "report": {
        "label": "Доклад",
        "title_page": False,
        "toc": False,
        "sections": [],
    },
    "free": {
        "label": "Только оформление",
        "title_page": False,
        "toc": False,
        "sections": [],
    },
}


# --- файлы и профили ---

def load_profile(path):
    if not os.path.isfile(path):
        raise ValueError("Файл профиля не найден.")
    try:
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    except json.JSONDecodeError:
        raise ValueError("Файл профиля содержит неверные данные JSON.")


def check_file_size(path, max_mb=MAX_DOCUMENT_SIZE_MB):
    size_mb = os.path.getsize(path) / (1024 * 1024)
    return size_mb <= max_mb, size_mb


def validate_docx_file(path):
    if not path:
        raise ValueError("Не указан путь к документу.")
    if not os.path.isfile(path):
        raise ValueError("Файл документа не найден.")
    if not path.lower().endswith(".docx"):
        raise ValueError("Можно выбрать только файл с расширением .docx.")
    return True


# --- работа с текстом ---

def clean(text):
    # "1. Введение:" -> "введение"
    text = text.strip().lower().rstrip(".:").strip()
    parts = text.split(" ", 1)
    if len(parts) == 2 and parts[0].rstrip(".)").replace(".", "").isdigit():
        text = parts[1].strip()
    return text


def is_subsection(text):
    # "1.1 Цели" или "2.3.1 Что-то"
    nums = text.strip().split(" ")[0].rstrip(".").split(".")
    return len(nums) > 1 and all(n.isdigit() for n in nums)


def split_dots(text):
    # "Введение.......3" -> ("Введение", "3"), если строка другая - None
    text = text.strip()
    i = len(text)
    while i > 0 and (text[i - 1].isdigit() or text[i - 1] in "-–"):
        i -= 1
    page, rest = text[i:], text[:i]
    if page == "":
        return None
    dots = 0
    while rest and rest[-1] in ".… \t":
        if rest[-1] in ".…":
            dots += 1
        rest = rest[:-1]
    if dots < 2 or rest.strip() == "":
        return None
    return rest.strip(), page


def split_space_page(text):
    # "Практическая часть 67" -> ("Практическая часть", "67"), если номер отделён пробелом
    text = text.strip()
    i = len(text)
    while i > 0 and (text[i - 1].isdigit() or text[i - 1] in "-–"):
        i -= 1
    page, rest = text[i:], text[:i]
    if page == "" or rest == "" or rest[-1] != " " or rest.strip() == "":
        return None
    return rest.strip(), page


def split_toc(text):
    text = text.strip()
    if "\t" in text:
        title, _, page = text.partition("\t")
        return title.strip(), page.strip()
    return split_dots(text) or (text, "")


def is_caption(text):
    t = text.strip().lower()
    return t.startswith(("рис.", "рисунок", "рис ", "рис("))


# --- работа с xml (в docx абзац это тег w:p, кусок текста - w:r) ---

def has_tag(el, name):
    for node in el.iter(qn(name)):
        return True
    return False


def get_runs(p):
    return [Run(r, p) for r in p._p.iter(qn("w:r"))]


def has_picture(p):
    return has_tag(p._p, "w:drawing") or has_tag(p._p, "w:pict")


def page_breaks(el):
    return [br for br in el.iter(qn("w:br")) if br.get(qn("w:type")) == "page"]


def break_before(p):
    ppr = p._p.pPr
    if ppr is None:
        return False
    flag = ppr.find(qn("w:pageBreakBefore"))
    return flag is not None and flag.get(qn("w:val")) not in ("0", "false")


def starts_with_break(p):
    # абзац начинается с ручного разрыва страницы
    for node in p._p.iter():
        if node.tag == qn("w:t") and (node.text or "").strip():
            return False
        if node.tag == qn("w:br") and node.get(qn("w:type")) == "page":
            return True
    return False


def get_sect(el):
    ppr = el.find(qn("w:pPr"))
    return None if ppr is None else ppr.find(qn("w:sectPr"))


def has_sect(el):
    return get_sect(el) is not None


def style_name(p):
    return (p.style.name or "").lower()


def has_num(p):
    ppr = p._p.pPr
    return ppr is not None and ppr.find(qn("w:numPr")) is not None


def is_list(p):
    return has_num(p) or style_name(p).startswith("list")


def is_toc_style(p):
    return style_name(p).startswith("toc")


def is_heading_style(p):
    return style_name(p).startswith(("heading", "заголовок"))


def is_top_heading(p):
    return style_name(p) in ("heading 1", "заголовок 1")


def before(p):
    # все абзацы документа перед этим, по порядку
    result = []
    for el in p._p.getroottree().getroot().iter(qn("w:p")):
        if el is p._p:
            break
        result.append(el)
    return result


def section_number(p):
    return len([el for el in before(p) if has_sect(el)])


def prev_par(p):
    node = p._p.getprevious()
    if node is not None and node.tag == qn("w:p"):
        return node
    return None


# --- чтение оформления (из текста или из стиля абзаца) ---

def from_style(p, get):
    style = p.style
    while style is not None:
        value = get(style)
        if value is not None:
            return value
        style = style.base_style
    return None


def get_font(run, p):
    try:
        return run.font.name or from_style(p, lambda s: s.font.name)
    except (ValueError, TypeError):
        return None


def get_size(run, p):
    try:
        size = run.font.size or from_style(p, lambda s: s.font.size)
    except (ValueError, TypeError):
        return None
    return size.pt if size else None


def get_bold(run, p):
    try:
        if run.font.bold is not None:
            return run.font.bold
        return bool(from_style(p, lambda s: s.font.bold))
    except (ValueError, TypeError):
        return False


def get_spacing(p):
    try:
        value = p.paragraph_format.line_spacing
        if value is None:
            value = from_style(p, lambda s: s.paragraph_format.line_spacing)
    except (ValueError, TypeError):
        return None
    # 1.5 приходит как float, а интервал в пунктах - другим типом
    return value if isinstance(value, float) else None


def get_indent(p):
    try:
        ind = p.paragraph_format.first_line_indent
        if ind is None:
            ind = from_style(p, lambda s: s.paragraph_format.first_line_indent)
    except (ValueError, TypeError):
        return 0.0
    return ind.cm if ind is not None else 0.0


# --- карта документа: титульный, оглавление, введение, темы ---

def find_toc_intro(pars):
    toc = intro = None
    for i, p in enumerate(pars):
        name = clean(p.text)
        if toc is None and name in TOC_NAMES:
            toc = i
        if name == "введение" and not is_toc_style(p) and (toc is None or i > toc):
            intro = i  # берем последнее, первое может быть в оглавлении
    return toc, intro


def find_title(pars, toc, intro):
    end = None
    for i, p in enumerate(pars[:TITLE_LIMIT + 1]):
        if i == toc or i == intro:
            end = i
        elif i > 0 and break_before(p):
            end = i
        elif page_breaks(p._p):
            end = i + 1 if p.text.strip() else i
        elif has_sect(p._p):
            end = i + 1
        if end is not None:
            break
    if end is None:
        return []
    title = pars[:end]
    if not any(p.text.strip() for p in title):
        return []
    return title


def find_toc_entries(pars, toc, intro):
    end = intro
    if end is None and toc is not None:
        # введения нет - оглавление кончается на первом разрыве страницы
        for i in range(toc + 1, len(pars)):
            if page_breaks(pars[i]._p) or break_before(pars[i]):
                end = i
                break
    result = []
    for i, p in enumerate(pars):
        if is_toc_style(p) or (toc is not None and end is not None and toc < i < end):
            result.append(p)
    return result


def get_toc_titles(entries):
    titles = []
    for p in entries:
        raw = split_toc(p.text)[0]
        name = clean(raw)
        if is_subsection(raw) or name == "" or name in titles:
            continue
        titles.append(name)
    return titles


def looks_like_heading(pars, i):
    # короткая строка, перед ней пустая строка, после нее длинный абзац
    p = pars[i]
    text = p.text.strip()
    if text == "" or "\n" in text or len(text) > 70 or len(text.split()) > 8:
        return False
    if text[0].islower() or text[0].isdigit() or text[0] in "-–—•":
        return False
    if text[-1] in ":,;" or is_caption(text) or has_picture(p) or is_list(p):
        return False
    if i == 0 or pars[i - 1].text.strip() != "":
        return False
    for nxt in pars[i + 1:]:
        if nxt.text.strip() != "":
            return len(nxt.text.strip()) >= 100
    return False


def find_themes(pars, title_el, toc_el, toc_titles, intro):
    themes = []
    found = []
    for p in pars:
        if p._p in title_el or p._p in toc_el:
            continue
        name = clean(p.text.split("\n")[0])  # заголовок может быть в одном абзаце с текстом
        if name == "":
            continue
        if is_top_heading(p) or name in KNOWN_TITLES or name in toc_titles:
            themes.append(p)
            found.append(name)

    # названия из оглавления, которых нет в тексте
    lost = [name for name in toc_titles if name not in found]
    if lost and intro is not None:
        stems = [w[:5] for title in lost for w in title.split() if len(w) >= 5]
        added = 0
        for i in range(intro + 1, len(pars)):
            if added >= len(lost):
                break
            p = pars[i]
            if p._p in toc_el or not looks_like_heading(pars, i):
                continue
            words = clean(p.text).split()
            if any(len(w) >= 5 and w[:5] in stems for w in words):
                themes.append(p)
                added += 1

    chosen = [p._p for p in themes]
    return [p for p in pars if p._p in chosen]


def analyze_layout(doc):
    pars = list(doc.paragraphs)
    toc, intro = find_toc_intro(pars)
    title = find_title(pars, toc, intro)
    toc_entries = find_toc_entries(pars, toc, intro)
    title_el = [p._p for p in title]
    toc_el = [p._p for p in toc_entries]
    themes = find_themes(pars, title_el, toc_el, get_toc_titles(toc_entries), intro)
    return {
        "title": title,
        "title_el": title_el,
        "toc": toc_entries,
        "toc_el": toc_el,
        "intro": pars[intro] if intro is not None else None,
        "themes": themes,
        "theme_el": [p._p for p in themes],
    }


def get_role(p, layout):
    # от роли абзаца зависит, какие правила к нему применять
    if p._p in layout["title_el"]:
        return "title"
    has_text = p.text.strip() != ""
    if not has_text and not has_picture(p):
        return "empty"
    if p._p in layout["toc_el"]:
        return "toc"
    if not has_text:
        return "image"
    if p._p in layout["theme_el"] or is_heading_style(p):
        return "heading"
    if is_list(p):
        return "list"
    if p.alignment in (WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT):
        return "center"
    return "body"


def table_pars(table):
    result = []
    for row in table.rows:
        for cell in row.cells:
            result += cell.paragraphs
            for inner in cell.tables:
                result += table_pars(inner)
    return result


def all_pars(doc, layout, with_title):
    result = [p for p in doc.paragraphs if with_title or p._p not in layout["title_el"]]
    for table in doc.tables:
        result += table_pars(table)
    return result


# --- проверки ---

def check_margins(doc, prof):
    want = [prof["margin_top_cm"], prof["margin_bottom_cm"], prof["margin_left_cm"], prof["margin_right_cm"]]
    for s in doc.sections:
        try:
            got = [s.top_margin.cm, s.bottom_margin.cm, s.left_margin.cm, s.right_margin.cm]
        except (ValueError, TypeError, AttributeError):
            return False
        for a, b in zip(got, want):
            if abs(a - b) > 0.05:
                return False
    return True


def check_font(doc, prof, layout):
    for p in all_pars(doc, layout, True):
        for run in get_runs(p):
            if run.text.strip() and get_font(run, p) != prof["font_name"]:
                return False
    return True


def check_size(doc, prof, layout):
    for p in all_pars(doc, layout, False):
        for run in get_runs(p):
            if run.text.strip():
                size = get_size(run, p)
                if size is None or abs(size - prof["font_size_pt"]) > 0.5:
                    return False
    return True


def check_spacing(doc, prof, layout):
    for p in doc.paragraphs:
        if get_role(p, layout) in SPACING_ROLES:
            s = get_spacing(p)
            if s is None or abs(s - prof["line_spacing"]) > 0.05:
                return False
    return True


def check_indent(doc, prof, layout):
    for p in doc.paragraphs:
        if get_role(p, layout) == "body":
            if abs(get_indent(p) - prof["first_line_indent_cm"]) > 0.1:
                return False
    return True


def structure_warnings(doc, prof):
    """Проверяет смысловую структуру, не пытаясь придумывать текст за автора."""
    document_type = prof.get("document_type", "free")
    rules = prof.get("structure") or DOCUMENT_TYPES.get(document_type, DOCUMENT_TYPES["free"])
    layout = analyze_layout(doc)
    warnings = []

    if rules.get("title_page") and not layout["title"]:
        warnings.append("В документе не найден обязательный титульный лист.")
    if rules.get("toc") and not layout["toc"]:
        warnings.append("В документе не найдено обязательное оглавление.")

    headings = {clean(p.text) for p in layout["themes"] if p.text.strip()}
    for section in rules.get("sections", []):
        if section not in headings:
            warnings.append("Не найден обязательный раздел «" + section.capitalize() + "».")
    return warnings


# номера страниц, напечатанные руками в тексте

def is_typed_num(p):
    text = p.text.strip()
    if text == "" or has_num(p):
        return False
    return text.isdigit() and text.isascii() and len(text) <= 3


def has_trailing_num(p):
    # "...текст.<Enter><табы>7"
    if is_typed_num(p):
        return False
    text = p.text.rstrip()
    if "\n" not in text:
        return False
    last = text.split("\n")[-1].strip()
    return last.isdigit() and last.isascii() and len(last) <= 3


def find_typed_nums(doc, layout):
    return [p for p in doc.paragraphs
            if p._p not in layout["title_el"] and (is_typed_num(p) or has_trailing_num(p))]


# номера в колонтитулах

def is_page_code(text):
    words = (text or "").split()
    return len(words) > 0 and words[0] == "PAGE"


def has_page_field(sections, i):
    # колонтитул "как в предыдущем разделе" - смотрим предыдущий
    while i > 0 and sections[i].footer.is_linked_to_previous:
        i -= 1
    footer = sections[i].footer
    if footer.is_linked_to_previous:
        return False
    for code in footer._element.iter(qn("w:instrText")):
        if is_page_code(code.text):
            return True
    for field in footer._element.iter(qn("w:fldSimple")):
        if is_page_code(field.get(qn("w:instr"))):
            return True
    return False


def check_numbers(doc, layout):
    sections = doc.sections
    intro = layout["intro"]
    if find_typed_nums(doc, layout):
        return False  # номера набраны руками, будет дубль

    if intro is None:
        return all(has_page_field(sections, i) for i in range(len(sections)))

    if intro._p.getprevious() is None:
        first = 0  # введение в самом начале
    else:
        prev = prev_par(intro)
        if prev is None or not has_sect(prev):
            return False  # перед введением нет разрыва раздела
        first = section_number(intro)

    for i, s in enumerate(sections):
        has_number = has_page_field(sections, i)
        if i < first and has_number:
            return False  # номер на титульном листе или оглавлении
        if i >= first:
            if not has_number or s.different_first_page_header_footer:
                return False
            if s._sectPr.find(qn("w:pgNumType")) is not None:
                return False  # нумерация начинается заново
    return True


# оглавление

def has_leader(p):
    return 'w:leader="dot"' in p._p.xml or 'w:leader="dot"' in p.style.element.xml


def check_toc(doc, layout):
    for p in layout["toc"]:
        text = p.text.strip()
        if text == "":
            continue
        if split_dots(text) is not None:
            return False  # точки набраны руками
        if "\t" not in text and split_space_page(text) is not None:
            return False  # номер страницы отделён пробелом, а не табуляцией
        if "\t" in text and not has_leader(p):
            return False
    return True


# подписи к картинкам

def caption_below(p):
    for el in p._p.itersiblings(qn("w:p")):
        nxt = Paragraph(el, p._parent)
        if nxt.text.strip() != "":
            return is_caption(nxt.text)
        if has_picture(nxt):
            return False
    return False


def has_caption(p):
    # подпись может быть строкой внутри абзаца с картинкой или в следующем абзаце
    if any(is_caption(line) for line in p.text.split("\n")):
        return True
    return caption_below(p)


def check_captions(doc, layout):
    for p in doc.paragraphs:
        if p._p not in layout["title_el"] and has_picture(p) and not has_caption(p):
            return False
    return True


# пунктуация

def punct_work(doc, layout, apply=False):
    # идёт по абзацам, находит правки; если apply=True - вносит их в документ
    changes = []
    biblio = False  # список литературы оформляется по-другому, его не трогаем
    prev_dot = None
    for p in doc.paragraphs:
        if p._p in layout["theme_el"] and clean(p.text.split("\n")[0]) in BIBLIO:
            biblio = True
        if biblio or is_typed_num(p):
            continue
        role = get_role(p, layout)
        nodes, text = get_nodes(p)
        long_heading = role == "heading" and len(text.strip()) > 100  # текст, которому дали стиль заголовка
        if role in ("body", "list") or long_heading:
            edits, prev_dot = find_edits(text, role == "list", not has_picture(p), prev_dot)
        elif p._p in layout["theme_el"]:
            edits = heading_edit(text)
            prev_dot = None
        else:
            edits = []
            prev_dot = None
        if apply:
            apply_edits(nodes, edits)
        changes += make_changes(text, edits)
    return changes


def check_punct(doc, layout):
    return punct_work(doc, layout) == []


# темы: с новой страницы и жирным

def new_page(p):
    if break_before(p) or starts_with_break(p):
        return True
    for el in reversed(before(p)):
        if has_sect(el) or page_breaks(el):
            return True
        prev = Paragraph(el, p._parent)
        if prev.text.strip() != "" or has_picture(prev):
            return False
    return True  # перед абзацем ничего нет


def title_runs(p):
    # если после названия в том же абзаце идет текст, берем куски до первого переноса строки
    runs = get_runs(p)
    if "\n" not in p.text:
        return runs
    result = []
    for run in runs:
        result.append(run)
        if run._r.find(qn("w:br")) is not None:
            break
    return result


def title_is_bold(p):
    return all(get_bold(run, p) for run in title_runs(p) if run.text.strip())


def check_themes(layout):
    for p in layout["themes"]:
        if not new_page(p) or not title_is_bold(p):
            return False
    return True


def check_document(doc, prof):
    # возвращает номера критериев, которые не прошли
    layout = analyze_layout(doc)
    results = {
        1: check_margins(doc, prof),
        2: check_font(doc, prof, layout),
        3: check_size(doc, prof, layout),
        4: check_spacing(doc, prof, layout),
        5: check_indent(doc, prof, layout),
        6: check_numbers(doc, layout),
        7: check_toc(doc, layout),
        8: check_captions(doc, layout),
        9: check_themes(layout),
        10: check_punct(doc, layout),
    }
    return [n for n in results if not results[n]]


def spell_warning(doc):
    # слова, которые Word подчеркнул красным; если пометок нет - ничего не пишем
    words = []
    for w in spell_words(doc):
        if w not in words:
            words.append(w)
    if not words:
        return []
    shown = ", ".join("«" + w + "»" for w in words[:10])
    more = " и ещё " + str(len(words) - 10) if len(words) > 10 else ""
    return ["Word подчеркнул красной линией слова: " + shown + more + ". Проверьте их правописание вручную."]


def find_warnings(doc):
    layout = analyze_layout(doc)
    warnings = []
    if not layout["title"]:
        warnings.append("Титульный лист не найден. Создайте его на вкладке «Титульный лист» и вставьте в документ.")
    if layout["intro"] is None:
        warnings.append("Раздел «Введение» не найден: номера страниц поставлены на все страницы, кроме первой.")
    return warnings


# --- исправления ---

def apply_margins(doc, prof):
    for s in doc.sections:
        s.top_margin = Cm(prof["margin_top_cm"])
        s.bottom_margin = Cm(prof["margin_bottom_cm"])
        s.left_margin = Cm(prof["margin_left_cm"])
        s.right_margin = Cm(prof["margin_right_cm"])


def apply_font(doc, prof):
    # сначала в стилях (заголовки берут шрифт из стиля), потом в тексте
    for style in doc.styles:
        if style.type == WD_STYLE_TYPE.PARAGRAPH:
            style.font.name = prof["font_name"]
            for attr in ["w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"]:
                style.element.rPr.rFonts.attrib.pop(qn(attr), None)  # шрифт темы перебивает обычный
    layout = analyze_layout(doc)
    for p in all_pars(doc, layout, True):
        for run in get_runs(p):
            run.font.name = prof["font_name"]


def apply_size(doc, prof):
    layout = analyze_layout(doc)
    for p in all_pars(doc, layout, False):
        for run in get_runs(p):
            run.font.size = Pt(prof["font_size_pt"])


def apply_spacing(doc, prof):
    layout = analyze_layout(doc)
    for p in doc.paragraphs:
        if get_role(p, layout) in SPACING_ROLES:
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
            p.paragraph_format.line_spacing = prof["line_spacing"]


def apply_indent(doc, prof):
    layout = analyze_layout(doc)
    for p in doc.paragraphs:
        role = get_role(p, layout)
        if role == "body":
            p.paragraph_format.first_line_indent = Cm(prof["first_line_indent_cm"])
        elif role in ("heading", "center"):
            p.paragraph_format.first_line_indent = Cm(0)


# нумерация: титульный лист и оглавление без номера, со введения - с номером

def strip_trailing_num(p):
    # убираем с конца абзаца число, табы и перенос строки перед ним
    nodes = [n for n in p._p.iter() if n.tag in (qn("w:t"), qn("w:br"), qn("w:tab"))]
    nodes.reverse()
    removed = False
    for node in nodes:
        if not removed:
            if node.tag != qn("w:t"):
                return
            node.getparent().remove(node)
            removed = True
        elif node.tag == qn("w:tab"):
            node.getparent().remove(node)
        elif node.tag == qn("w:br") and node.get(qn("w:type")) != "page":
            node.getparent().remove(node)
            return
        else:
            return


def remove_typed_nums(doc):
    layout = analyze_layout(doc)
    for p in find_typed_nums(doc, layout):
        if has_trailing_num(p):
            strip_trailing_num(p)
        elif has_sect(p._p) or page_breaks(p._p):
            for t in p._p.iter(qn("w:t")):
                t.text = ""  # разрыв оставляем, число убираем
        else:
            p._p.getparent().remove(p._p)


def add_page_field(p, prof):
    # номер страницы в word - поле из 4 частей: begin, PAGE, separate, end
    for kind in ["begin", "code", "separate", "end"]:
        run = p.add_run()
        run.font.name = prof["font_name"]
        run.font.size = Pt(prof["font_size_pt"])
        if kind == "code":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = " PAGE "
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        run._r.append(el)


def clear_footer(footer):
    # удаляем все (и w:p, и w:sdt, который вставляет word) и оставляем один пустой абзац
    el = footer._element
    for child in list(el):
        el.remove(child)
    new_p = OxmlElement("w:p")
    el.append(new_p)
    return Paragraph(new_p, footer)


def write_number(footer, prof):
    p = clear_footer(footer)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_field(p, prof)


def find_sect_props(p):
    # настройки раздела, в котором лежит абзац
    own = get_sect(p._p)
    if own is not None:
        return own
    for el in p._p.itersiblings(qn("w:p")):
        found = get_sect(el)
        if found is not None:
            return found
    return p._p.getparent().find(qn("w:sectPr"))


def remove_breaks_before(p):
    # ручные разрывы страницы перед абзацем заменим разрывом раздела
    for el in p._p.itersiblings(preceding=True):
        if el.tag != qn("w:p"):
            break
        for br in page_breaks(el):
            br.getparent().remove(br)
        wrapped = Paragraph(el, p._parent)
        if wrapped.text.strip() != "" or has_picture(wrapped):
            break


def add_section_break(p):
    if p._p.getprevious() is None:
        return
    prev = prev_par(p)
    if prev is not None and has_sect(prev):
        return  # разрыв уже есть

    remove_breaks_before(p)
    props = find_sect_props(p)
    new_props = copy.deepcopy(props)
    # первый раздел не должен ссылаться на колонтитулы второго
    for tag in ["w:headerReference", "w:footerReference", "w:titlePg", "w:type"]:
        for el in new_props.findall(qn(tag)):
            new_props.remove(el)
    for el in props.findall(qn("w:type")):
        props.remove(el)  # раздел со введением всегда с новой страницы

    if prev is None:
        prev = OxmlElement("w:p")
        p._p.addprevious(prev)
    prev.get_or_add_pPr().append(new_props)


def fix_numbers(doc, prof):
    remove_typed_nums(doc)
    layout = analyze_layout(doc)

    if layout["intro"] is None:
        # введения нет - нумеруем все, кроме первой страницы
        for i, s in enumerate(doc.sections):
            s.footer.is_linked_to_previous = False
            write_number(s.footer, prof)
            s.different_first_page_header_footer = (i == 0)
            if i == 0:
                s.first_page_footer.is_linked_to_previous = False
                clear_footer(s.first_page_footer)
        return

    add_section_break(layout["intro"])
    first = section_number(layout["intro"])
    for i, s in enumerate(doc.sections):
        s.footer.is_linked_to_previous = False
        if i < first:
            clear_footer(s.footer)
        else:
            write_number(s.footer, prof)
            s.different_first_page_header_footer = False
            for el in s._sectPr.findall(qn("w:pgNumType")):
                s._sectPr.remove(el)  # убираем "начать нумерацию заново"


# оглавление

def dots_to_tab(p):
    # "Введение.....3" -> "Введение<таб>3", чтобы номера встали ровно
    if "\t" in p.text:
        return
    result = split_dots(p.text) or split_space_page(p.text)
    if result is None or p._p.find(qn("w:hyperlink")) is not None:
        return
    runs = p.runs
    if not runs:
        return
    runs[0].text = result[0] + "\t" + result[1]
    for run in runs[1:]:
        run._r.getparent().remove(run._r)


def fix_toc(doc):
    layout = analyze_layout(doc)
    s = doc.sections[0]
    width = s.page_width.cm - s.left_margin.cm - s.right_margin.cm
    for p in layout["toc"]:
        if p.text.strip() == "":
            continue
        dots_to_tab(p)
        # У оглавления не должно быть старого абзацного отступа:
        # иначе правый табулятор выходит за границу строки и номер страницы переносится.
        p.paragraph_format.left_indent = Cm(0)
        p.paragraph_format.right_indent = Cm(0)
        p.paragraph_format.first_line_indent = Cm(0)
        if not has_leader(p):
            p.paragraph_format.tab_stops.add_tab_stop(Cm(width), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)


# подписи к картинкам

def add_captions(doc, prof):
    layout = analyze_layout(doc)
    number = 0
    for p in list(doc.paragraphs):
        if p._p in layout["title_el"] or not has_picture(p):
            continue
        number += 1
        if has_caption(p):
            continue

        el = OxmlElement("w:p")
        p._p.addnext(el)
        cap = Paragraph(el, p._parent)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        cap.paragraph_format.line_spacing = prof["line_spacing"]
        cap.paragraph_format.first_line_indent = Cm(0)
        run = cap.add_run("Рис. " + str(number))
        run.font.name = prof["font_name"]
        run.font.size = Pt(prof["font_size_pt"])

        if p.text.strip() == "":
            # центрируем только абзац с одной картинкой, абзац с текстом не трогаем
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True


# темы: с новой страницы и жирным

def fix_themes(doc):
    layout = analyze_layout(doc)
    for p in layout["themes"]:
        if not new_page(p):
            p.paragraph_format.page_break_before = True
        for run in title_runs(p):
            run.font.bold = True


# --- титульный лист ---

def set_run_style(p, bold, size, upper=False):
    for run in p.runs:
        run.font.bold = bold
        run.font.size = Pt(size)
        if upper:
            run.text = run.text.upper()


def format_title(doc):
    # по зонам: справа - участники, по центру - название работы и город, остальное жирным
    layout = analyze_layout(doc)
    center = []
    for p in layout["title"]:
        if p.text.strip() == "":
            continue
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.paragraph_format.first_line_indent = Cm(0)
        if p.alignment == WD_ALIGN_PARAGRAPH.RIGHT:
            set_run_style(p, False, 14)
        elif p.alignment == WD_ALIGN_PARAGRAPH.CENTER:
            center.append(p)
        else:
            set_run_style(p, True, 14)
    if len(center) >= 1:
        set_run_style(center[0], True, 16, upper=True)
    if len(center) >= 2:
        set_run_style(center[-1], True, 14)


def insert_title_page(doc, title_path):
    validate_docx_file(title_path)
    body = doc.element.body
    new_els = [copy.deepcopy(el) for el in Document(title_path).element.body if el.tag != qn("w:sectPr")]

    # пустые абзацы в начале (место под титульный) убираем, чтобы не было пустой страницы
    while len(body) > 1 and body[0].tag == qn("w:p"):
        first = Paragraph(body[0], doc._body)
        if first.text.strip() != "" or has_picture(first) or has_sect(body[0]):
            break
        body.remove(body[0])

    for i, el in enumerate(new_els):
        body.insert(i, el)

    # текст документа начинаем с новой страницы
    first_old = body[len(new_els)]
    target = None
    if first_old.tag == qn("w:p"):
        target = first_old
    else:
        for node in first_old.iter(qn("w:p")):
            target = node
            break
    if target is not None:
        Paragraph(target, doc._body).paragraph_format.page_break_before = True
    else:
        br_par = OxmlElement("w:p")
        run = OxmlElement("w:r")
        br = OxmlElement("w:br")
        br.set(qn("w:type"), "page")
        run.append(br)
        br_par.append(run)
        first_old.addprevious(br_par)


# --- главная функция ---

def check_and_format_document(input_path, profile_path, output_path, title_path=None):
    # проверяем документ; если все хорошо - не трогаем, иначе исправляем нарушенное
    validate_docx_file(input_path)
    prof = load_profile(profile_path)

    doc = Document(input_path)
    failed = check_document(doc, prof)
    had_title = len(analyze_layout(doc)["title"]) > 0
    spell = spell_warning(doc)

    if not failed and not title_path:
        return {"is_gost": True, "failed": [], "title_inserted": False, "changes": [],
                "warnings": find_warnings(doc) + structure_warnings(doc, prof) + spell}

    doc = Document(input_path)  # чистая копия
    changes = []
    if 10 in failed:
        # пунктуацию правим до вставки титульного листа, пока старый титульный лист виден как титульный
        changes = punct_work(doc, analyze_layout(doc), apply=True)
    if title_path:
        insert_title_page(doc, title_path)
        # после вставки структура изменилась - проверяем заново
        for n in check_document(doc, prof):
            if n not in failed and n != 10:
                failed.append(n)
        failed.sort()

    # сначала структура, потом оформление
    if 8 in failed:
        add_captions(doc, prof)
    if 6 in failed:
        fix_numbers(doc, prof)
    if 9 in failed:
        fix_themes(doc)
    if 1 in failed:
        apply_margins(doc, prof)
    if 2 in failed:
        apply_font(doc, prof)
    if 3 in failed:
        apply_size(doc, prof)
    if 4 in failed:
        apply_spacing(doc, prof)
    if 5 in failed:
        apply_indent(doc, prof)
    if 7 in failed:
        fix_toc(doc)
    if any(n in failed for n in (2, 3, 4, 5)):
        format_title(doc)

    doc.save(output_path)

    warnings = find_warnings(doc) + structure_warnings(doc, prof)
    if title_path and had_title:
        warnings.insert(0, "В документе уже был свой титульный лист — теперь их два. Удалите лишний.")
    return {"is_gost": False, "failed": failed, "title_inserted": bool(title_path),
            "changes": changes, "warnings": warnings + spell}


# --- профили ---

def list_profiles(folder="profiles"):
    if not os.path.isdir(folder):
        return []
    profiles = []
    for name in sorted(os.listdir(folder)):
        path = os.path.join(folder, name)
        if name.lower().endswith(".json") and os.path.isfile(path):
            profiles.append({"file": name, "name": load_profile(path).get("name", "Без названия")})
    return profiles


def create_profile(
    name,
    font_name,
    font_size_pt,
    line_spacing,
    first_line_indent_cm,
    margin_top_cm,
    margin_bottom_cm,
    margin_left_cm,
    margin_right_cm,
    document_type="free",
    profiles_folder="profiles",
):
    if not name:
        raise ValueError("Укажите название профиля.")
    if not font_name:
        raise ValueError("Укажите название шрифта.")
    if document_type not in DOCUMENT_TYPES:
        raise ValueError("Выбран неизвестный тип документа.")

    os.makedirs(profiles_folder, exist_ok=True)
    profile = {
        "name": name,
        "font_name": font_name,
        "font_size_pt": font_size_pt,
        "line_spacing": line_spacing,
        "first_line_indent_cm": first_line_indent_cm,
        "margin_top_cm": margin_top_cm,
        "margin_bottom_cm": margin_bottom_cm,
        "margin_left_cm": margin_left_cm,
        "margin_right_cm": margin_right_cm,
        "document_type": document_type,
        "structure": copy.deepcopy(DOCUMENT_TYPES[document_type]),
    }
    number = 1
    while os.path.exists(os.path.join(profiles_folder, "profile_" + str(number) + ".json")):
        number += 1
    path = os.path.join(profiles_folder, "profile_" + str(number) + ".json")
    with open(path, "w", encoding="utf-8") as f:
        json.dump(profile, f, ensure_ascii=False, indent=4)
    return path
